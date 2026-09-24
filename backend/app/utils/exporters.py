"""论文导出工具：Markdown → DOCX / PDF

- DOCX：python-docx，正文宋体 12pt、1.5 倍行距、首行缩进 2 字符，标题黑体居中
- PDF：reportlab，使用内置 CJK 字体 STSong-Light（无需额外字体文件）
"""

import io
import re
from typing import Any, Dict

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import Paragraph, SimpleDocTemplate
from xml.sax.saxutils import escape

from app.utils.helpers import strip_ai_meta

# 中文字体
_BODY_FONT = "宋体"
_HEADING_FONT = "黑体"
# PDF 内置中文 CID 字体
_PDF_FONT = "STSong-Light"
_pdf_font_registered = False

# Markdown 行内标记
_BOLD_PATTERN = re.compile(r"\*\*(.+?)\*\*")
_INLINE_CODE_PATTERN = re.compile(r"`(.+?)`")
_HEADING_PATTERN = re.compile(r"^(#{1,6})\s+(.*)$")
_BULLET_PATTERN = re.compile(r"^([-*•]|\d+[.、)])\s+(.*)$")


def _clean_inline(text: str) -> str:
    """去掉 Markdown 行内标记"""
    text = _BOLD_PATTERN.sub(r"\1", text)
    text = _INLINE_CODE_PATTERN.sub(r"\1", text)
    return re.sub(r"^>\s*", "", text).strip()


# ----------------------------------------------------------------------
# DOCX
# ----------------------------------------------------------------------
def _set_run_font(run, name: str, size: int = 12, bold: bool = False) -> None:
    """设置中英文字体（东亚字体需单独设置 w:eastAsia）"""
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)


def markdown_to_docx(title: str, content: str) -> bytes:
    """把论文 Markdown 转换为 DOCX 字节流"""
    document = Document()

    normal = document.styles["Normal"]
    normal.font.name = _BODY_FONT
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), _BODY_FONT)

    # 论文标题
    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(heading.add_run(title or "论文"), _HEADING_FONT, size=18, bold=True)

    for raw_line in strip_ai_meta(content).split("\n"):
        line = _clean_inline(raw_line)
        if not line:
            continue

        heading_match = _HEADING_PATTERN.match(raw_line.strip())
        if heading_match:
            level = len(heading_match.group(1))
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(10 if level <= 2 else 6)
            paragraph.paragraph_format.space_after = Pt(6)
            if level == 1:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _set_run_font(
                paragraph.add_run(_clean_inline(heading_match.group(2))),
                _HEADING_FONT,
                size={1: 16, 2: 14, 3: 13}.get(level, 12),
                bold=True,
            )
            continue

        bullet_match = _BULLET_PATTERN.match(raw_line.strip())
        if bullet_match:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Pt(24)
            paragraph.paragraph_format.space_after = Pt(4)
            marker = "" if bullet_match.group(1)[0].isdigit() else "• "
            _set_run_font(
                paragraph.add_run(marker + _clean_inline(bullet_match.group(2))),
                _BODY_FONT,
            )
            continue

        paragraph = document.add_paragraph()
        paragraph.paragraph_format.first_line_indent = Pt(24)
        paragraph.paragraph_format.line_spacing = 1.5
        paragraph.paragraph_format.space_after = Pt(6)
        _set_run_font(paragraph.add_run(line), _BODY_FONT)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# ----------------------------------------------------------------------
# PDF
# ----------------------------------------------------------------------
def _ensure_pdf_font() -> None:
    """注册 reportlab 内置中文 CID 字体（无需外部字体文件）"""
    global _pdf_font_registered
    if not _pdf_font_registered:
        pdfmetrics.registerFont(UnicodeCIDFont(_PDF_FONT))
        _pdf_font_registered = True


def _pdf_styles() -> Dict[Any, ParagraphStyle]:
    """PDF 段落样式"""
    return {
        1: ParagraphStyle("h1", fontName=_PDF_FONT, fontSize=18, leading=28, alignment=1, spaceAfter=12),
        2: ParagraphStyle("h2", fontName=_PDF_FONT, fontSize=15, leading=24, spaceBefore=12, spaceAfter=8),
        3: ParagraphStyle("h3", fontName=_PDF_FONT, fontSize=13, leading=22, spaceBefore=8, spaceAfter=6),
        "body": ParagraphStyle("body", fontName=_PDF_FONT, fontSize=12, leading=22, firstLineIndent=24, spaceAfter=6),
        "bullet": ParagraphStyle("bullet", fontName=_PDF_FONT, fontSize=12, leading=22, leftIndent=24, bulletIndent=10, spaceAfter=6),
    }


def markdown_to_pdf(title: str, content: str) -> bytes:
    """把论文 Markdown 转换为 PDF 字节流"""
    _ensure_pdf_font()
    styles = _pdf_styles()

    story = [Paragraph(escape(title or "论文"), styles[1])]
    for raw_line in strip_ai_meta(content).split("\n"):
        stripped = raw_line.strip()
        if not stripped:
            continue

        heading_match = _HEADING_PATTERN.match(stripped)
        if heading_match:
            level = len(heading_match.group(1))
            story.append(Paragraph(escape(_clean_inline(heading_match.group(2))), styles.get(level, styles[3])))
            continue

        bullet_match = _BULLET_PATTERN.match(stripped)
        if bullet_match:
            story.append(Paragraph(
                escape(_clean_inline(bullet_match.group(2))),
                styles["bullet"],
                bulletText="•" if not bullet_match.group(1)[0].isdigit() else bullet_match.group(1),
            ))
            continue

        story.append(Paragraph(escape(_clean_inline(stripped)), styles["body"]))

    buffer = io.BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=2.5 * cm,
        rightMargin=2.5 * cm,
        topMargin=2.5 * cm,
        bottomMargin=2.5 * cm,
        title=title or "论文",
    )
    document.build(story)
    return buffer.getvalue()